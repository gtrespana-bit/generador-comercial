"""Capa imperial de cantidades (EE. UU., Fase 1).

El almacenamiento es siempre métrico; en el mercado imperial la superficie
(editor, fichas, PDF, exportaciones) trabaja en ft / sq ft / cu yd / gal /
lb, con la cantidad multiplicada por el factor y el precio dividido por él,
de modo que el importe de cada línea es idéntico en los dos sistemas.

La marca que hace seguro el viaje de vuelta es la propia unidad: un payload
que llega con unidad «sq ft» trae sus cantidades en pies cuadrados y el
servidor lo normaliza a m² al guardar, sin banderas en el formulario.

Estas pruebas fijan:

* la invariancia del importe en la ida y la vuelta;
* que los precios de recurso del APU NO se convierten (solo el rendimiento,
  que es «por unidad de partida») y que la fila «%» no trata su rendimiento
  como una cantidad;
* que un presupuesto estadounidense se crea en sq ft, se guarda en m² y se
  vuelve a abrir en sq ft, y que el PDF y el detalle hablan en pies;
* que el catálogo (ficha, «actualizar precio», guardar desde el editor) y
  los packs quedan en métrico en la base aunque se editen en imperial;
* que los otros mercados no cambian en absoluto.
"""

import io
import json
from datetime import date

import pytest
from pypdf import PdfReader

from app.models import (
    Capitulo,
    Cliente,
    Configuracion,
    Partida,
    Presupuesto,
    PresupuestoItem,
    PresupuestoItemProducto,
    RecetaEstancia,
)
from app.services import unidades_precio as u

F_M2 = 10.76391042


# ─────────────────────────────────────────────────────────────────────────────
# Servicio: factores, etiquetas y valores sueltos
# ─────────────────────────────────────────────────────────────────────────────

def test_factores_y_etiquetas_de_la_capa():
    assert u.factor_cantidad("m") == pytest.approx(3.280839895)
    assert u.factor_cantidad("m2") == pytest.approx(F_M2)
    assert u.factor_cantidad("m²") == pytest.approx(F_M2)
    assert u.factor_cantidad("ml") == pytest.approx(3.280839895)
    assert u.factor_cantidad("m3") == pytest.approx(1.307950619)
    assert u.factor_cantidad("kg") == pytest.approx(2.204622622)
    assert u.factor_cantidad("l") == pytest.approx(0.2641720524)
    assert u.factor_cantidad("ud") is None
    assert u.factor_cantidad("glb") is None
    assert u.factor_cantidad("h") is None
    assert u.factor_cantidad("") is None

    assert u.unidad_mostrada("m2", True) == "sq ft"
    assert u.unidad_mostrada("m3", True) == "cu yd"
    assert u.unidad_mostrada("m", True) == "ft"
    assert u.unidad_mostrada("l", True) == "gal"
    assert u.unidad_mostrada("kg", True) == "lb"
    assert u.unidad_mostrada("ud", True) == "ud"
    assert u.unidad_mostrada("m2", False) == "m2"


def test_el_importe_es_invariante_en_la_conversion():
    for unidad, cantidad, precio in (
        ("m2", 12.5, 50.0), ("m", 30.0, 8.4), ("m3", 4.2, 209.27),
        ("kg", 850.0, 1.32), ("l", 23.0, 10.57), ("ud", 3.0, 99.0),
    ):
        cant_i = u.cantidad_mostrada(cantidad, unidad, True)
        precio_i = u.precio_mostrado(precio, unidad, True)
        assert cant_i * precio_i == pytest.approx(cantidad * precio, rel=1e-3)

    # El rendimiento se mueve como el precio (es «por unidad de partida»).
    assert u.rendimiento_mostrado(0.85, "m2", True) == pytest.approx(0.85 / F_M2, rel=1e-5)
    # Sin mercado imperial, nada se toca.
    assert u.cantidad_mostrada(12.5, "m2", False) == 12.5
    assert u.precio_mostrado(50.0, "m2", False) == 50.0
    assert u.precio_mostrado(None, "m2", True) is None


def test_la_puerta_es_el_pais_de_la_empresa():
    class Cfg:
        def __init__(self, pais):
            self.empresa_pais = pais

    assert u.es_imperial_config(Cfg("Estados Unidos")) is True
    assert u.es_imperial_config(Cfg("United States")) is True
    assert u.es_imperial_config(Cfg("España")) is False
    assert u.es_imperial_config(Cfg("Reino Unido")) is False
    assert u.es_imperial_config(Cfg("")) is False
    assert u.es_imperial_config(None) is False


# ─────────────────────────────────────────────────────────────────────────────
# Servicio: normalizar el payload del editor y volver a enseñarlo
# ─────────────────────────────────────────────────────────────────────────────

def _partida_editor_imperial() -> dict:
    """Una partida tal como sale del parser cuando el editor va en sq ft."""
    return {
        "unidad": "sq ft",
        "cantidad": 134.5489,            # 12,5 m²
        "precio": 4.6452,                # 50 $/m²
        "coste_materiales": 2.0,
        "coste_mano_obra": 1.5,
        "coste_complementarios": 0.1,
        "coste_otros": 0.0,
        "prod_precio": 0.5,
        "prod_coste": 0.3,
        "prod_unidad": "sq ft",
        "mediciones": [("Salón", 100.0), ("Cocina", 34.5489)],
        "productos_opciones": [
            {"precio": 1.0, "coste": 0.6, "unidad": "sq ft"},
            {"precio": 25.0, "coste": None, "unidad": "ud"},
        ],
        "descomposicion": [
            {"tipo": "recurso", "unidad": "kg", "rendimiento": 0.464515,
             "precio": 0.446, "importe": 0.2072},
            {"tipo": "recurso", "unidad": "%", "rendimiento": 3,
             "precio": 3.6, "importe": 0.108},
            {"tipo": "subtotal", "unidad": "", "rendimiento": "",
             "precio": "", "importe": 3.6},
        ],
        "descomposicion_meta": {"unidad": "sq ft"},
    }


def test_normalizar_partida_a_metrico_completa():
    pd = _partida_editor_imperial()
    importe_antes = pd["cantidad"] * pd["precio"]

    assert u.normalizar_partida_a_metrico(pd) is True

    assert pd["unidad"] == "m2"
    assert pd["cantidad"] == pytest.approx(12.5, abs=1e-3)
    assert pd["precio"] == pytest.approx(50.0, abs=1e-2)
    assert pd["cantidad"] * pd["precio"] == pytest.approx(importe_antes, rel=1e-4)
    assert pd["coste_materiales"] == pytest.approx(2.0 * F_M2)
    assert pd["prod_precio"] == pytest.approx(0.5 * F_M2)
    assert pd["prod_unidad"] == "m2"
    assert pd["mediciones"][0][1] == pytest.approx(100.0 / F_M2)
    # Las opciones se suman por unidad de partida: convierten con su factor
    # aunque su etiqueta diga «ud»; la etiqueta solo acompaña si coincidía.
    assert pd["productos_opciones"][0]["precio"] == pytest.approx(F_M2)
    assert pd["productos_opciones"][0]["unidad"] == "m2"
    assert pd["productos_opciones"][1]["precio"] == pytest.approx(25.0 * F_M2)
    assert pd["productos_opciones"][1]["unidad"] == "ud"
    # APU: el rendimiento convierte; el precio del recurso ($/kg) no.
    fila = pd["descomposicion"][0]
    assert fila["rendimiento"] == pytest.approx(5.0, abs=1e-3)
    assert fila["precio"] == 0.446
    assert fila["importe"] == pytest.approx(0.2072 * F_M2)
    # Fila %: el rendimiento es un porcentaje (intacto) y su «precio» es la
    # base monetaria por unidad de partida (convierte).
    fila_pct = pd["descomposicion"][1]
    assert fila_pct["rendimiento"] == 3
    assert fila_pct["precio"] == pytest.approx(3.6 * F_M2)
    assert pd["descomposicion"][2]["importe"] == pytest.approx(3.6 * F_M2)
    assert pd["descomposicion_meta"]["unidad"] == "m2"


def test_normalizar_no_toca_lo_que_no_es_imperial():
    for unidad in ("m2", "ud", "glb", "lote", ""):
        pd = {"unidad": unidad, "cantidad": 10.0, "precio": 5.0,
              "mediciones": [], "productos_opciones": [], "descomposicion": []}
        assert u.normalizar_partida_a_metrico(pd) is False
        assert pd["cantidad"] == 10.0 and pd["precio"] == 5.0


def test_la_vuelta_a_pantalla_desde_lo_guardado():
    pd = _partida_editor_imperial()
    u.normalizar_partida_a_metrico(pd)
    vista = u.partida_editor_a_mostrada(pd, True)

    assert vista["unidad"] == "sq ft"
    assert vista["cantidad"] == pytest.approx(134.5489, abs=1e-2)
    assert vista["precio"] == pytest.approx(4.6452, abs=1e-3)
    assert vista["descomposicion"][0]["rendimiento"] == pytest.approx(0.464515, abs=1e-4)
    assert vista["descomposicion"][1]["rendimiento"] == 3
    # Sin mercado imperial la vista es el propio dict.
    assert u.partida_editor_a_mostrada(pd, False) is pd


def test_ficha_del_catalogo_a_pantalla():
    ficha = {
        "unidad": "m2", "precio": 50.0, "precio_unitario": 50.0,
        "coste_materiales": 21.5278, "tiempo_estimado_horas": 0.85,
        "descomposicion": {"unidad": "m2", "filas": [
            {"tipo": "recurso", "unidad": "h", "rendimiento": 0.85,
             "precio": 39.07, "importe": 33.2095},
            {"tipo": "recurso", "unidad": "%", "rendimiento": 3,
             "precio": 38.75, "precio_unitario": 38.75, "importe": 1.1625},
        ]},
    }
    vista = u.ficha_catalogo_a_mostrada(ficha, True)
    assert vista["unidad"] == "sq ft"
    assert vista["precio"] == pytest.approx(50.0 / F_M2, rel=1e-4)
    assert vista["coste_materiales"] == pytest.approx(2.0, abs=1e-3)
    # Las horas por unidad se mueven como un rendimiento.
    assert vista["tiempo_estimado_horas"] == pytest.approx(0.85 / F_M2, rel=1e-5)
    f0 = vista["descomposicion"]["filas"][0]
    assert f0["rendimiento"] == pytest.approx(0.85 / F_M2, rel=1e-5)
    assert f0["precio"] == 39.07  # $/h intacto
    f1 = vista["descomposicion"]["filas"][1]
    assert f1["precio"] == pytest.approx(38.75 / F_M2, rel=1e-4)
    assert f1["precio_unitario"] == f1["precio"]
    # La ficha original no se muta.
    assert ficha["precio"] == 50.0 and ficha["unidad"] == "m2"


def test_pack_ida_y_vuelta_con_coeficientes():
    receta = {
        "unidad_base": "m²", "cantidad_base_default": 10.0,
        "items": [{
            "nombre": "Alicatado", "unidad": "m2", "precio": 50.0,
            "tipo_calculo": "proporcional", "coeficiente": 1.2,
            "cantidad_fija": 12.0, "prod_precio": 5.0, "prod_unidad": "m2",
        }],
    }
    vista = u.receta_a_mostrada(receta, True)
    assert vista["unidad_base"] == "sq ft"
    assert vista["cantidad_base_default"] == pytest.approx(10.0 * F_M2, rel=1e-4)
    item = vista["items"][0]
    assert item["unidad"] == "sq ft"
    assert item["precio"] == pytest.approx(50.0 / F_M2, rel=1e-4)
    # Ítem y base en la misma unidad: el coeficiente no cambia.
    assert item["coeficiente"] == pytest.approx(1.2)

    # Guardar un capítulo medido en sq ft deja el pack en métrico.
    items = [{"nombre": "Alicatado", "unidad": "sq ft", "cantidad": 129.167,
              "precio": 4.6452, "prod_precio": 0.4645, "prod_unidad": "sq ft"}]
    base, cantidad_base = u.normalizar_pack_a_metrico(items, "sq ft", 107.639)
    assert base == "m2"
    assert cantidad_base == pytest.approx(10.0, abs=1e-2)
    assert items[0]["unidad"] == "m2"
    assert items[0]["cantidad"] == pytest.approx(12.0, abs=1e-2)
    assert items[0]["precio"] == pytest.approx(50.0, abs=1e-1)


# ─────────────────────────────────────────────────────────────────────────────
# PDF (servicio, sin HTTP): el documento del cliente habla en pies
# ─────────────────────────────────────────────────────────────────────────────

def _presupuesto_pdf():
    cliente = Cliente(nombre="Client", rif="X-1")
    p = Presupuesto(
        numero="P-US-1", titulo="Kitchen remodel", fecha=date(2026, 10, 1),
        moneda="USD", estado="borrador", impuesto_pct=7.0, cliente=cliente,
    )
    cap = Capitulo(nombre="FLOORING", orden=1)
    partida = PresupuestoItem(
        nombre="Porcelain tile flooring", unidad="m2",
        cantidad=12.5, precio_unitario=50.0,
    )
    partida.productos_opciones = [
        PresupuestoItemProducto(nombre="Tile A", precio=12.0, unidad="m2", seleccionado=True),
        PresupuestoItemProducto(nombre="Tile B", precio=30.0, unidad="m2", seleccionado=False),
    ]
    cap.partidas = [partida]
    p.capitulos = [cap]
    return p, partida


def _texto_pdf(buf) -> str:
    lector = PdfReader(io.BytesIO(buf.getvalue()))
    return "\n".join((pagina.extract_text() or "") for pagina in lector.pages)


def test_pdf_en_pies_para_empresa_estadounidense_y_en_metros_para_espanola():
    from app.services.pdf import generar_pdf

    p_us, _ = _presupuesto_pdf()
    cfg_us = Configuracion(empresa_nombre="Sunshine Builders LLC",
                            empresa_pais="Estados Unidos", pdf_color="#0F4C81")
    buf_us = generar_pdf(p_us, cfg_us, lang="en")
    texto_us = _texto_pdf(buf_us)
    assert "sq ft" in texto_us
    assert "134.55" in texto_us          # 12,5 m² en pies cuadrados
    assert "4,645" not in texto_us       # el precio va en formato US y por sq ft
    # La partida tiene dos productos a elegir → PDF interactivo: el importe
    # vive en un campo AcroForm. El P.U. se cotiza a céntimos por sq ft
    # (4,65 $/sq ft × 134,55 = 625,65), como multiplicaría un estimate en
    # papel; el documento es coherente consigo mismo al céntimo.
    campos = PdfReader(io.BytesIO(buf_us.getvalue())).get_fields() or {}
    valores = {k: v.get("/V") for k, v in campos.items()}
    assert valores.get("pu_p1") == "4.65 USD/sq ft"
    assert valores.get("imp_p1") == "625.65 USD"

    p_es, _ = _presupuesto_pdf()
    cfg_es = Configuracion(empresa_nombre="Reformas Sol",
                            empresa_pais="España", pdf_color="#0F4C81")
    texto_es = _texto_pdf(generar_pdf(p_es, cfg_es, lang="es"))
    assert "sq ft" not in texto_es
    assert "m2" in texto_es


def test_contexto_interactivo_convierte_su_tabla_y_conserva_importes():
    from app.services.pdf_interactivo import ContextoInteractivo

    p, partida = _presupuesto_pdf()
    ctx = ContextoInteractivo(p, "USD", None, None, lang="en", imperial=True)
    assert ctx.preparar() is True
    pid = ctx.id_partida(partida)
    datos = ctx.partidas[pid]
    assert datos["cantidad"] == pytest.approx(12.5 * F_M2, rel=1e-4)
    assert datos["precios"][0] == pytest.approx(12.0 / F_M2, rel=1e-3)
    assert datos["unidades"][0] == "sq ft"
    # El importe recalculado (Python y JS comparten la tabla) es el métrico
    # salvo el redondeo del P.U. a céntimos por sq ft: 4,65 × 134,5489.
    assert ctx.pu(pid) == pytest.approx(4.65)
    assert ctx.imp(pid) == pytest.approx(625.65, abs=0.01)
    assert ctx.imp(pid) == pytest.approx(12.5 * 50.0, rel=2e-3)


def test_excel_del_documento_en_pies_y_catalogo_en_metros():
    from openpyxl import load_workbook

    from app.services.excel_export import exportar_presupuesto_excel

    p, _ = _presupuesto_pdf()
    cfg = Configuracion(empresa_nombre="Sunshine Builders LLC",
                        empresa_pais="Estados Unidos")
    buf = exportar_presupuesto_excel(p, cfg, idioma="en")
    wb = load_workbook(io.BytesIO(buf.getvalue()))
    ws = wb["Detail"] if "Detail" in wb.sheetnames else wb[wb.sheetnames[1]]
    celdas = [c.value for fila in ws.iter_rows() for c in fila]
    assert "sq ft" in celdas
    assert any(isinstance(v, (int, float)) and abs(v - 12.5 * F_M2) < 0.1 for v in celdas)


def test_ejemplo_de_la_landing_en_pies_para_eeuu():
    from app.services.landing_ejemplo import contexto_ejemplo

    ctx_us = contexto_ejemplo("US", "en")
    cants_us = [f["cant"] for cap in ctx_us["caps"] for f in cap["partidas"]]
    assert any("sq ft" in c for c in cants_us)

    ctx_es = contexto_ejemplo("ES", "es")
    cants_es = [f["cant"] for cap in ctx_es["caps"] for f in cap["partidas"]]
    assert not any("sq ft" in c for c in cants_es)
    # Los importes del ejemplo no dependen del sistema de medida: el mismo
    # ejemplo en un mercado USD métrico (Ecuador) suma lo mismo que en EE. UU.
    # (solo difiere el impuesto), y el APU pasa de $/m² a $/sq ft.
    ctx_ec = contexto_ejemplo("EC", "es")
    assert ctx_us["tot"]["subtotal"] == ctx_ec["tot"]["subtotal"]
    assert ctx_us["tot"]["obra"] == ctx_ec["tot"]["obra"]
    assert ctx_us["apu"]["unidad"] == "sq ft"
    assert "/sq ft" in ctx_us["apu"]["precio"]
    assert ctx_ec["apu"]["unidad"] == "m²"


# ─────────────────────────────────────────────────────────────────────────────
# HTTP: crear en sq ft, guardar en m², reabrir en sq ft
# ─────────────────────────────────────────────────────────────────────────────

def _empresa_eeuu(Session):
    with Session() as db:
        cfg = db.query(Configuracion).first()
        cfg.empresa_pais = "Estados Unidos"
        cfg.moneda_default = "USD"
        db.commit()
        cliente = db.query(Cliente).first()
        return cliente.id


def _estructura_sqft():
    return [{
        "nombre": "FLOORING",
        "partidas": [{
            "nombre": "Porcelain tile",
            "unidad": "sq ft",
            "cantidad": 134.55,
            "precio": 4.65,
            "descripcion": "",
            "mediciones": [{"concepto": "Kitchen", "cantidad": 134.55}],
            "descomposicion": [
                {"tipo": "recurso", "grupo": "Mano de obra", "categoria": "mano_obra",
                 "codigo": "MO-OF1", "unidad": "h", "descripcion": "Oficial",
                 "rendimiento": 0.079, "precio": 39.07, "importe": 3.09},
            ],
        }],
    }]


def test_crear_en_sq_ft_guarda_en_m2_y_reabre_en_sq_ft(entorno, cliente_web):
    Session, _ids, _rol = entorno
    cliente_id = _empresa_eeuu(Session)

    respuesta = cliente_web.post(
        "/presupuestos/nuevo",
        data={
            "client_id": str(cliente_id), "titulo": "Kitchen",
            "fecha": date.today().isoformat(), "validez_dias": "30",
            "moneda": "USD", "impuesto_pct": "7", "descuento_pct": "0",
            "estado": "borrador",
            "estructura_json": json.dumps(_estructura_sqft()),
        },
        follow_redirects=False,
    )
    assert respuesta.status_code == 303
    pid = int(respuesta.headers["location"].split("?")[0].split("/")[-1])

    with Session() as db:
        partida = (
            db.query(PresupuestoItem)
            .join(Capitulo)
            .filter(Capitulo.presupuesto_id == pid)
            .one()
        )
        assert partida.unidad == "m2"
        assert partida.cantidad_total == pytest.approx(134.55 / F_M2, rel=1e-3)
        assert partida.precio_unitario == pytest.approx(4.65 * F_M2, rel=1e-3)
        assert partida.importe == pytest.approx(134.55 * 4.65, abs=0.05)
        fila = partida.descomposicion_cype.filas[0]
        assert fila.rendimiento == pytest.approx(0.079 * F_M2, rel=1e-3)
        assert fila.precio_unitario == pytest.approx(39.07)

    # El editor lo vuelve a enseñar en pies (datos iniciales convertidos).
    html_editar = cliente_web.get(f"/presupuestos/{pid}/editar").text
    assert '"unidad": "sq ft"' in html_editar
    assert "134.55" in html_editar

    # El detalle y el PDF también.
    html_detalle = cliente_web.get(f"/presupuestos/{pid}").text
    assert "sq ft" in html_detalle
    pdf = cliente_web.get(f"/presupuestos/{pid}/pdf")
    assert pdf.status_code == 200
    lector = PdfReader(io.BytesIO(pdf.content))
    texto = "\n".join((pg.extract_text() or "") for pg in lector.pages)
    assert "sq ft" in texto


def test_la_unidad_es_la_marca_tambien_fuera_de_eeuu(entorno, cliente_web):
    """Un payload en «lb» se guarda en kg venga de donde venga: la unidad
    dice en qué están las cantidades y el almacenamiento es métrico."""
    Session, _ids, _rol = entorno
    with Session() as db:
        cfg = db.query(Configuracion).first()
        cfg.empresa_pais = "España"
        db.commit()
        cliente_id = db.query(Cliente).first().id

    estructura = [{
        "nombre": "ESTRUCTURA",
        "partidas": [{
            "nombre": "Acero", "unidad": "lb", "cantidad": 2204.62,
            "precio": 0.6, "descripcion": "", "mediciones": [],
        }],
    }]
    respuesta = cliente_web.post(
        "/presupuestos/nuevo",
        data={
            "client_id": str(cliente_id), "titulo": "Acero",
            "fecha": date.today().isoformat(), "validez_dias": "30",
            "moneda": "USD", "impuesto_pct": "21", "descuento_pct": "0",
            "estado": "borrador", "estructura_json": json.dumps(estructura),
        },
        follow_redirects=False,
    )
    assert respuesta.status_code == 303
    pid = int(respuesta.headers["location"].split("?")[0].split("/")[-1])
    with Session() as db:
        partida = (
            db.query(PresupuestoItem).join(Capitulo)
            .filter(Capitulo.presupuesto_id == pid).one()
        )
        assert partida.unidad == "kg"
        assert partida.cantidad == pytest.approx(1000.0, rel=1e-3)
        assert partida.precio_unitario == pytest.approx(0.6 * 2.204622622, rel=1e-3)


def test_ficha_y_actualizar_precio_del_catalogo_en_pies(entorno, cliente_web):
    Session, _ids, _rol = entorno
    _empresa_eeuu(Session)
    with Session() as db:
        partida = db.query(Partida).first()
        partida.unidad = "m2"
        partida.precio_unitario = 50.0
        db.commit()
        partida_id = partida.id

    ficha = cliente_web.get(f"/partidas/{partida_id}/ficha").json()
    assert ficha["partida"]["unidad"] == "sq ft"
    assert ficha["partida"]["precio"] == pytest.approx(50.0 / F_M2, rel=1e-3)

    # El editor envía el nuevo precio en $/sq ft; el catálogo guarda $/m².
    respuesta = cliente_web.post(
        f"/partidas/{partida_id}/actualizar-precio",
        json={"precio": 5.0, "moneda": "USD"},
    ).json()
    assert respuesta["ok"] is True
    assert respuesta["partida"]["precio"] == pytest.approx(5.0, rel=1e-3)
    with Session() as db:
        assert db.get(Partida, partida_id).precio_unitario == pytest.approx(5.0 * F_M2, rel=1e-3)


def test_guardar_ficha_desde_presupuesto_en_sq_ft_persiste_en_metros(entorno, cliente_web):
    Session, _ids, _rol = entorno
    _empresa_eeuu(Session)
    with Session() as db:
        partida_id = db.query(Partida).first().id

    respuesta = cliente_web.post(
        "/partidas/guardar-desde-presupuesto",
        data={
            "partida_catalogo_id": str(partida_id),
            "nombre": "Drywall partition (imperial test)",
            "unidad": "sq ft",
            "categoria": "Drywall",
            "precio_unitario": "4.6452",
            "coste_materiales": "2.0",
            "coste_mano_obra": "1.5",
            "coste_complementarios": "0.1",
            "coste_otros": "0",
            "moneda": "USD",
            "d_categoria": ["mano_obra"],
            "d_codigo": ["MO-OF1"],
            "d_unidad": ["h"],
            "d_descripcion": ["Oficial"],
            "d_rendimiento": ["0.079"],
            "d_precio": ["39.07"],
        },
    ).json()
    assert respuesta["ok"] is True
    assert respuesta["partida"]["unidad"] == "sq ft"
    with Session() as db:
        partida = db.get(Partida, partida_id)
        assert partida.unidad == "m2"
        assert partida.precio_unitario == pytest.approx(50.0, abs=0.05)
        filas = json.loads(partida.descomposicion_json)["filas"]
        assert filas[0]["rendimiento"] == pytest.approx(0.079 * F_M2, rel=1e-3)
        assert filas[0]["precio"] == pytest.approx(39.07)


def test_pack_guardado_en_sq_ft_se_lee_en_sq_ft_y_se_guarda_en_metros(entorno, cliente_web):
    Session, _ids, _rol = entorno
    _empresa_eeuu(Session)

    creado = cliente_web.post(
        "/recetas/api/guardar-desde-capitulo",
        json={
            "nombre": "Bathroom tile (imperial test)",
            "unidad_base": "sq ft",
            "cantidad_base_default": 107.639,
            "moneda": "USD",
            "items": [{"nombre": "Tile work", "unidad": "sq ft",
                       "cantidad": 107.639, "precio": 4.6452}],
        },
    ).json()
    assert creado["ok"] is True

    with Session() as db:
        receta = db.get(RecetaEstancia, creado["id"])
        assert receta.unidad_base == "m2"
        assert receta.cantidad_base_default == pytest.approx(10.0, abs=0.01)
        item = json.loads(receta.datos)[0]
        assert item["unidad"] == "m2"
        assert item["precio"] == pytest.approx(50.0, abs=0.05)
        assert item["coeficiente"] == pytest.approx(1.0)

    listado = cliente_web.get("/recetas/api/list").json()
    pack = next(r for r in listado["recetas"] if r["id"] == creado["id"])
    assert pack["unidad_base"] == "sq ft"
    assert pack["items"][0]["unidad"] == "sq ft"
    assert pack["items"][0]["precio"] == pytest.approx(4.6452, rel=1e-3)


def test_plantilla_guardada_en_sq_ft_se_lee_en_sq_ft(entorno, cliente_web):
    Session, _ids, _rol = entorno
    _empresa_eeuu(Session)

    estructura = _estructura_sqft()
    creada = cliente_web.post(
        "/plantillas",
        data={"nombre": "Kitchen tile (imperial test)", "datos": json.dumps(estructura)},
    ).json()
    assert creada["ok"] is True

    datos = cliente_web.get(f"/plantillas/{creada['id']}/datos").json()
    partida = datos["capitulos"][0]["partidas"][0]
    assert partida["unidad"] == "sq ft"
    assert partida["cantidad"] == pytest.approx(134.55, rel=1e-3)
    assert partida["precio"] == pytest.approx(4.65, rel=1e-3)

    with Session() as db:
        from app.models import Plantilla

        guardada = json.loads(db.query(Plantilla).filter_by(id=creada["id"]).one().datos)
        assert guardada[0]["partidas"][0]["unidad"] == "m2"


def test_mercado_metrico_sin_cambios_en_la_ficha(entorno, cliente_web):
    """Regresión: fuera de EE. UU. la ficha sale exactamente como siempre."""
    Session, _ids, _rol = entorno
    with Session() as db:
        cfg = db.query(Configuracion).first()
        cfg.empresa_pais = "España"
        partida = db.query(Partida).first()
        partida.unidad = "m2"
        partida.precio_unitario = 50.0
        db.commit()
        partida_id = partida.id

    ficha = cliente_web.get(f"/partidas/{partida_id}/ficha").json()
    assert ficha["partida"]["unidad"] == "m2"
    assert ficha["partida"]["precio"] == pytest.approx(50.0)
