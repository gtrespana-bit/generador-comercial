"""Excel y CSV de exportación en inglés.

El libro se lo lleva un cliente: si el panel está en inglés, las hojas, las
cabeceras y los rótulos también. Lo que no cambia son los números: un
presupuesto exportado en los dos idiomas tiene que cuadrar al céntimo.
"""
import io
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.services import excel_export as X


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------

def _partida(nombre="Porcelanato", **kw):
    datos = dict(
        nombre=nombre, descripcion="", unidad="m2", cantidad_total=2.0,
        precio_unitario=50.0, importe=100.0, coste_materiales=10.0,
        coste_mano_obra=5.0, coste_complementarios=1.0, coste_otros=0.0,
        desperdicio_pct=5.0, tipo_partida="included", seleccionada=True,
        coste_directo=None, producto=None, producto_id=None,
    )
    datos.update(kw)
    return SimpleNamespace(**datos)


def _capitulo(nombre="PISOS", partidas=None):
    return SimpleNamespace(nombre=nombre, partidas=partidas or [_partida()])


def _presupuesto(**kw):
    p = SimpleNamespace(
        numero="P-2026-001", year=2026, fecha=None, estado="borrador",
        titulo="Reforma de baño", direccion_obra="C/ Mayor 1",
        codigo_postal="28001", cliente=SimpleNamespace(nombre="Familia R"),
        moneda="USD", tipo_cambio=None, fecha_tipo_cambio=None,
        impuesto_pct=16.0, descuento_pct=0.0, validez_dias=30,
        notas="Nota", condiciones="Condición", usar_funciones_avanzadas=False,
        capitulos=[_capitulo()], subtotal=100.0, subtotal_opcional=0.0,
        subtotal_alternativas=0.0, total_productos=0.0, subtotal_obra=100.0,
        costes_adicionales=0.0, descuento_monto=0.0, base=100.0,
        impuesto_monto=16.0, total=116.0, coste_interno=16.0,
        margen=84.0, margen_pct=72.4, margen_obra=84.0, margen_obra_pct=72.4,
        margen_productos=0.0, margen_productos_pct=0.0,
    )
    p.todas_partidas = [i for c in p.capitulos for i in c.partidas]
    for k, v in kw.items():
        setattr(p, k, v)
    return p


def _hojas(buf):
    wb = load_workbook(io.BytesIO(buf.getvalue()))
    return {ws.title: [list(f) for f in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


# ---------------------------------------------------------------------------
# Libro del presupuesto
# ---------------------------------------------------------------------------

def test_libro_en_castellano_no_cambia():
    libro = _hojas(X.exportar_presupuesto_excel(_presupuesto()))
    assert list(libro) == ["Resumen", "Detalle", "Costes Internos"]
    assert libro["Resumen"][0][0] == "RESUMEN DEL PRESUPUESTO"
    assert libro["Detalle"][2][:3] == ["Capítulo", "Partida", "Descripción"]


def test_libro_en_ingles_traduce_hojas_y_cabeceras():
    libro = _hojas(X.exportar_presupuesto_excel(_presupuesto(), idioma="en"))
    assert list(libro) == ["Summary", "Line items", "Internal costs"]
    assert libro["Summary"][0][0] == "ESTIMATE SUMMARY"
    assert libro["Line items"][2][:3] == ["Section", "Line item", "Description"]
    assert libro["Internal costs"][2][:3] == ["Section", "Line item", "Material cost"]


def test_libro_en_ingles_traduce_los_rotulos_del_resumen():
    resumen = _hojas(X.exportar_presupuesto_excel(_presupuesto(), idioma="en"))["Summary"]
    etiquetas = [f[0] for f in resumen if f and isinstance(f[0], str)]
    for esperado in ("Document details", "Number", "Client", "Postal code",
                     "VAT (%)", "Validity (days)", "TOTALS", "TOTAL",
                     "Additional information", "Commercial terms"):
        assert esperado in etiquetas, f"falta {esperado!r}"
    for prohibido in ("Datos del documento", "Número", "Cliente", "Código postal",
                      "IVA (%)", "Validez (días)", "TOTALES",
                      "Información adicional", "Condiciones comerciales"):
        assert prohibido not in etiquetas, f"queda {prohibido!r}"


def test_los_numeros_no_cambian_con_el_idioma():
    """El idioma cambia las palabras, nunca los importes."""
    p = _presupuesto()
    es = _hojas(X.exportar_presupuesto_excel(p))
    en = _hojas(X.exportar_presupuesto_excel(p, idioma="en"))
    for titulo_es, titulo_en in zip(es, en):
        filas_es, filas_en = es[titulo_es], en[titulo_en]
        assert len(filas_es) == len(filas_en)
        for a, b in zip(filas_es, filas_en):
            assert [c for c in a if not isinstance(c, str)] == \
                   [c for c in b if not isinstance(c, str)]


def test_tabla_de_costes_solo_aparece_si_hay_costes():
    p = _presupuesto()
    assert "Costes Internos" in _hojas(X.exportar_presupuesto_excel(p))
    sin_costes = _partida(coste_materiales=0.0, coste_mano_obra=0.0,
                          coste_complementarios=0.0, coste_otros=0.0)
    p2 = _presupuesto(capitulos=[_capitulo(partidas=[sin_costes])])
    p2.todas_partidas = [sin_costes]
    assert "Internal costs" not in _hojas(X.exportar_presupuesto_excel(p2, idioma="en"))


def test_sin_capitulo_y_subtotales_en_ingles():
    p = _presupuesto(capitulos=[_capitulo(nombre=None)])
    detalle = _hojas(X.exportar_presupuesto_excel(p, idioma="en"))["Line items"]
    celdas = [c for f in detalle for c in f if isinstance(c, str)]
    assert "No section" in celdas
    assert any(c.startswith("Subtotal No section") for c in celdas)
    assert "TOTAL INCLUDED" in celdas


# ---------------------------------------------------------------------------
# Historial y catálogos
# ---------------------------------------------------------------------------

def test_historial_en_ingles():
    p = _presupuesto(fecha=None)
    es = _hojas(X.exportar_historial_excel([p]))
    en = _hojas(X.exportar_historial_excel([p], idioma="en"))
    assert list(es) == ["Presupuestos"]
    assert list(en) == ["Estimates"]
    assert es["Presupuestos"][3][:4] == ["Nº", "Fecha", "Cliente", "Título"]
    assert en["Estimates"][3][:4] == ["No.", "Date", "Client", "Title"]
    assert en["Estimates"][0][0] == "Estimate history"


def test_catalogo_de_partidas_en_ingles():
    from app.models import Partida

    p = Partida(nombre="Porcelanato", unidad="m2", precio_unitario=12.0,
                categoria="Albañilería", descripcion="")
    es = _hojas(X.exportar_catalogo_partidas_excel([p]))
    en = _hojas(X.exportar_catalogo_partidas_excel([p], idioma="en"))
    assert list(es) == ["Partidas"] and list(en) == ["Line items"]
    assert es["Partidas"][3][:3] == ["Código", "Código anterior", "Nombre"]
    assert en["Line items"][3][:3] == ["Code", "Previous code", "Name"]
    assert "Labour cost" in en["Line items"][3]
    assert "Coste mano obra" not in en["Line items"][3]


def test_catalogo_de_productos_y_recursos_en_ingles():
    prod = SimpleNamespace(
        nombre="Cuarzo", marca="X", modelo="Y", sku="Z", descripcion="",
        unidad="m2", precio_compra=1.0, precio_unitario=2.0, categoria="C",
        proveedor="P", color="Blanco", acabado="Pulido", formato="60x120",
        tiempo_entrega_dias=3, variantes="", ficha_tecnica="",
        fecha_actualizacion_precio=None, usos=0,
    )
    en = _hojas(X.exportar_catalogo_productos_excel([prod], idioma="en"))
    assert list(en) == ["Products"]
    assert en["Products"][3][:4] == ["Name", "Brand", "Model", "SKU"]
    assert "Lead time (days)" in en["Products"][3]

    rec = SimpleNamespace(codigo="R1", descripcion="Oficial", unidad="h",
                          categoria="mano_obra", grupo="G", id=1, usos=0,
                          proveedor="", fecha_actualizacion_precio=None)
    en = _hojas(X.exportar_catalogo_recursos_excel([rec], idioma="en", moneda="USD"))
    assert list(en) == ["Resources"]
    assert en["Resources"][3][:5] == ["Code", "Description", "Unit", "Category", "Group"]
    assert en["Resources"][3][5] == "Price (USD)"


def test_total_de_cada_catalogo_se_cuenta_en_ingles():
    from app.models import Partida

    p = Partida(nombre="Porcelanato", unidad="m2", precio_unitario=12.0)
    en = _hojas(X.exportar_catalogo_partidas_excel([p, p], idioma="en"))
    assert en["Line items"][1][0] == "Total: 2 line items"


# ---------------------------------------------------------------------------
# Cabeceras compartidas con los CSV de los routers
# ---------------------------------------------------------------------------

def test_cabeceras_de_csv_comparten_vocabulario_con_el_excel():
    assert X.cabeceras(["Número", "Base", "Total"], "es") == ["Número", "Base", "Total"]
    assert X.cabeceras(["Número", "Base", "Total"], "en") == ["Number", "Taxable base", "Total"]
    # Variante regional: manda la lengua.
    assert X.cabeceras(["Descripción"], "en-mx") == ["Description"]
    assert X.cabeceras(["Última actualización de precio"], "en") == ["Price last updated"]
    # Una etiqueta sin traducir se devuelve igual, nunca vacía.
    assert X.cabeceras(["Rareza"], "en") == ["Rareza"]


def test_toda_etiqueta_traducible_tiene_su_par_y_no_se_repite():
    # TOTAL, Subtotal y SKU son iguales en los dos idiomas.
    _iguales = {"TOTAL", "Total", "Subtotal", "SKU"}
    for es, en in X._XLSX_EN.items():
        assert es and en, f"etiqueta vacía: {es!r} -> {en!r}"
        if es in _iguales:
            continue
        assert es != en, f"traducción igual al original: {es!r}"


# ---------------------------------------------------------------------------
# Rutas HTTP
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("ruta,esperado,prohibido", [
    ("/recursos/exportar", "Code", "Código"),
    ("/partidas/exportar", "Unit price", "Precio unitario"),
    ("/productos/exportar", "Brand", "Marca"),
])
def test_csv_de_catalogo_sigue_el_idioma_del_panel(entorno, cliente_web, ruta, esperado, prohibido):
    cliente_web.cookies.set("cotizat_lang", "en")
    csv = cliente_web.get(ruta).content.decode("utf-8")
    assert esperado in csv
    assert prohibido not in csv.split("\r\n")[0]

    cliente_web.cookies.set("cotizat_lang", "es")
    csv_es = cliente_web.get(ruta).content.decode("utf-8")
    assert prohibido in csv_es.split("\r\n")[0]
    cliente_web.cookies.clear()


def test_csv_del_presupuesto_sigue_el_idioma_del_panel(entorno, cliente_web):
    Session, _ids, _rol = entorno
    with Session() as db:
        from app.models import Presupuesto
        p = db.query(Presupuesto).order_by(Presupuesto.id).first()
        pid, numero = p.id, p.numero

    cliente_web.cookies.set("cotizat_lang", "en")
    csv = cliente_web.get(f"/presupuestos/{pid}/exportar").content.decode("utf-8")
    assert "Section" in csv and "Amount" in csv
    assert "TAXABLE BASE" in csv and "TOTAL" in csv

    cliente_web.cookies.set("cotizat_lang", "es")
    csv_es = cliente_web.get(f"/presupuestos/{pid}/exportar").content.decode("utf-8")
    assert "Capítulo" in csv_es and "BASE IMPONIBLE" in csv_es
    cliente_web.cookies.clear()
    # Los importes no dependen del idioma.
    assert "160,00" in csv and "160,00" in csv_es
    assert "185,60" in csv and "185,60" in csv_es
    assert numero not in csv  # el número viaja en el nombre del archivo
